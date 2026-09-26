import docx
import pytest
from pypdf import PdfWriter

from core.importers import (
    SUPPORTED_EXTENSIONS,
    TranscriptImportError,
    import_docx,
    import_file,
    import_pdf,
    import_srt,
    import_txt,
)


def _write_text_pdf(path, text_line):
    """Writes a small single-page PDF that really contains the given text,
    so it's not a mock. pypdf can read and modify PDFs but can't write text
    into one, so this builds the raw PDF structure by hand, with just enough
    for PdfReader.extract_text() to find the text."""
    content_stream = f"BT /F1 24 Tf 72 700 Td ({text_line}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content_stream)).encode() + b" >>\nstream\n"
        + content_stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode()
    path.write_bytes(bytes(out))


# .txt files

def test_import_txt_plain_utf8(tmp_path):
    path = tmp_path / "plain.txt"
    path.write_bytes("Hello world. Café résumé.\n".encode("utf-8"))
    assert import_txt(str(path)) == "Hello world. Café résumé.\n"

def test_import_txt_strips_utf8_bom(tmp_path):
    path = tmp_path / "bom.txt"
    path.write_bytes("Hello with BOM.\n".encode("utf-8-sig"))
    assert import_txt(str(path)) == "Hello with BOM.\n"

def test_import_txt_falls_back_to_latin1(tmp_path):
    path = tmp_path / "latin1.txt"
    path.write_bytes("Café in Latin-1.\n".encode("latin-1"))
    assert import_txt(str(path)) == "Café in Latin-1.\n"

def test_import_file_rejects_empty_txt(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_bytes(b"")
    with pytest.raises(TranscriptImportError, match="readable text"):
        import_file(str(path))


# .srt files

def test_import_srt_drops_index_and_timestamp_lines(tmp_path):
    path = tmp_path / "sample.srt"
    path.write_text(
        "1\n00:00:01,000 --> 00:00:04,000\nHello world.\n\n"
        "2\n00:00:04,500 --> 00:00:06,000\nSecond line.\n",
        encoding="utf-8",
    )
    result = import_srt(str(path))
    assert "-->" not in result
    assert "Hello world." in result
    assert "Second line." in result
    assert "1" not in result.split()
    assert "2" not in result.split()

def test_import_srt_strips_simple_markup_tags(tmp_path):
    path = tmp_path / "italic.srt"
    path.write_text("1\n00:00:01,000 --> 00:00:02,000\n<i>Emphasized</i> text.\n", encoding="utf-8")
    assert import_srt(str(path)) == "Emphasized text."


# .docx files

def test_import_docx_extracts_paragraphs_and_drops_blank_ones(tmp_path):
    path = tmp_path / "sample.docx"
    d = docx.Document()
    d.add_paragraph("First paragraph.")
    d.add_paragraph("")
    d.add_paragraph("Second paragraph.")
    d.save(path)
    assert import_docx(str(path)) == "First paragraph.\nSecond paragraph."

def test_import_docx_rejects_corrupt_file(tmp_path):
    path = tmp_path / "corrupt.docx"
    path.write_bytes(b"not a real docx file")
    with pytest.raises(TranscriptImportError):
        import_docx(str(path))

def test_import_file_rejects_empty_docx(tmp_path):
    path = tmp_path / "empty.docx"
    docx.Document().save(path)
    with pytest.raises(TranscriptImportError, match="readable text"):
        import_file(str(path))


# .pdf files

def test_import_pdf_extracts_real_text(tmp_path):
    path = tmp_path / "sample.pdf"
    _write_text_pdf(path, "Hello from a real PDF")
    assert "Hello from a real PDF" in import_pdf(str(path))

def test_import_pdf_rejects_corrupt_file(tmp_path):
    path = tmp_path / "corrupt.pdf"
    path.write_bytes(b"%PDF-1.4 not really valid")
    with pytest.raises(TranscriptImportError):
        import_pdf(str(path))

def test_import_file_rejects_pdf_with_no_extractable_text(tmp_path):
    path = tmp_path / "no_text.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(path, "wb") as f:
        writer.write(f)
    with pytest.raises(TranscriptImportError, match="readable text"):
        import_file(str(path))

def test_import_pdf_decrypts_when_user_password_is_empty(tmp_path):
    """Some PDFs are encrypted with only an owner password and an empty user
    password. This is common for files that just block printing or editing
    but open without a password. pypdf's is_encrypted stays True even after
    a successful decrypt(), so import_pdf() has to check what decrypt()
    returns instead. This test makes sure that path really works, not just
    that it doesn't crash."""
    path = tmp_path / "encrypted_emptyuserpw.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(user_password="", owner_password="ownersecret")
    with open(path, "wb") as f:
        writer.write(f)
    # The blank page has no text, so this should fail with "no readable
    # text" and not "password-protected". That proves the decrypt with the
    # empty password succeeded first.
    with pytest.raises(TranscriptImportError, match="readable text"):
        import_file(str(path))

def test_import_pdf_reports_password_protected_when_password_is_required(tmp_path):
    path = tmp_path / "encrypted_realpw.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(user_password="realpassword", owner_password="ownersecret2")
    with open(path, "wb") as f:
        writer.write(f)
    with pytest.raises(TranscriptImportError, match="password-protected"):
        import_file(str(path))


# Picking the importer by file extension

def test_import_file_dispatches_by_extension(tmp_path):
    path = tmp_path / "plain.txt"
    path.write_text("Hello world.", encoding="utf-8")
    assert import_file(str(path)) == "Hello world."

def test_import_file_rejects_unsupported_extension(tmp_path):
    path = tmp_path / "sample.xyz"
    path.write_text("irrelevant", encoding="utf-8")
    with pytest.raises(TranscriptImportError, match="Unsupported"):
        import_file(str(path))

def test_supported_extensions_contents():
    assert SUPPORTED_EXTENSIONS == {".txt", ".srt", ".docx", ".pdf"}